"""
Builds the final EIP report workbook from a raw calendar export JSON,
reusing the deterministic filter logic in filter_eip_events.py.

Usage:
    python scripts/export_report.py <input_export.json> <output.xlsx>
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from filter_eip_events import load_events, filter_and_sort  # noqa: E402

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ["Date", "Start", "End", "Event Title", "Category Label(s)",
           "EIP Flag", "Notes / Description", "Attendees", "Organizer"]


def build_workbook(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "EIP Calendar Export"

    bold = Font(bold=True, name="Arial")
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    eip_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    wrap = Alignment(wrap_text=True, vertical="top")

    for col, h in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = bold
        c.fill = header_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")

    for r, row in enumerate(rows, 2):
        values = [row["date"], row["start_time"], row["end_time"], row["title"],
                  row["categories"], row["eip_flag"], row["notes"],
                  row["attendees"], row["organizer"]]
        for c, val in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = wrap
            if row["eip_flag"] == "Yes":
                cell.fill = eip_fill

    widths = [12, 10, 10, 42, 16, 8, 55, 40, 28]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{len(rows) + 1}"
    return wb


def main():
    if len(sys.argv) < 3:
        print("Usage: python export_report.py <input_export.json> <output.xlsx>", file=sys.stderr)
        sys.exit(1)

    input_path, output_path = sys.argv[1], sys.argv[2]
    events = load_events(input_path)
    rows = filter_and_sort(events)
    wb = build_workbook(rows)
    wb.save(output_path)
    eip_count = sum(1 for r in rows if r["eip_flag"] == "Yes")
    print(f"Wrote {len(rows)} events ({eip_count} EIP-flagged) to {output_path}")


if __name__ == "__main__":
    main()
